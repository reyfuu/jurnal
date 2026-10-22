# /// script
# dependencies = ["openpyxl"]
# ///
"""Rekap status unduhan peraturan ke Excel.

Jalankan: uv run scripts/peraturan/rekap-excel.py
Hasil:    docs/peraturan/rekap_download.xlsx
Status dihitung dari keberadaan berkas di disk, bukan dari kolom index.csv.
"""
import collections
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2] / "docs" / "peraturan"
HIJAU = PatternFill("solid", fgColor="C6EFCE")
KUNING = PatternFill("solid", fgColor="FFEB9C")
MERAH = PatternFill("solid", fgColor="FFC7CE")
HEAD = PatternFill("solid", fgColor="1F2937")


def status(row):
    f = row["file"]
    if f and f != "tanpa-pdf" and (ROOT / f).is_file():
        return "Sudah (HTML)" if f.endswith((".htm", ".html")) else "Sudah (PDF)"
    return "Belum"


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(1, i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def isi(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    cols = ["No", "Status Download", "Kategori", "Kode", "Judul", "Jenis", "Tahun",
            "Status Hukum", "Berkas", "Ukuran (KB)", "Catatan", "Primbon", "JDIH"]
    header(ws, cols, [6, 16, 14, 30, 60, 20, 7, 12, 36, 11, 40, 18, 18])
    for n, r in enumerate(rows, 1):
        s = r["_status"]
        kb = round(int(r["ukuran_byte"]) / 1024) if r["ukuran_byte"].isdigit() else None
        ws.append([n, s, r["kategori"], r["kode"], r["judul"], r["jenis"], r["tahun"],
                   r["status"], r["file"] if s != "Belum" else "", kb, r["catatan"],
                   r["primbon_url"], r["jdih_url"]])
        ws.cell(n + 1, 2).fill = MERAH if s == "Belum" else KUNING if "HTML" in s else HIJAU
        for col, url in ((12, r["primbon_url"]), (13, r["jdih_url"])):
            if url:
                c = ws.cell(n + 1, col)
                c.hyperlink, c.value, c.style = url, "buka", "Hyperlink"
    ws.auto_filter.ref = ws.dimensions

    ring = wb.create_sheet("Ringkasan", 0)
    header(ring, ["Kategori", "Total", "Sudah (PDF)", "Sudah (HTML)", "Belum", "% Terunduh"],
           [22, 9, 12, 13, 9, 12])
    per = collections.defaultdict(collections.Counter)
    for r in rows:
        per[r["kategori"]][r["_status"]] += 1
    tot = collections.Counter()
    for k in sorted(per):
        c = per[k]
        tot.update(c)
        ring.append(baris(k, c))
    ring.append(baris("TOTAL", tot))
    for cell in ring[ring.max_row]:
        cell.font = Font(bold=True)
    for row in ring.iter_rows(min_row=2):
        row[5].number_format = "0.0%"

    belum = wb.create_sheet("Belum Terdownload")
    header(belum, ["No", "Kategori", "Kode", "Judul", "Alasan", "Primbon"], [6, 14, 30, 60, 50, 60])
    for n, r in enumerate((r for r in rows if r["_status"] == "Belum"), 1):
        belum.append([n, r["kategori"], r["kode"], r["judul"],
                      r["catatan"] or r["sumber_ln"] or "PDF tidak tersedia di JDIH", r["primbon_url"]])
    for ws_ in wb.worksheets:
        for row in ws_.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top")
    return wb, tot


def baris(nama, c):
    total = sum(c.values())
    sudah = c["Sudah (PDF)"] + c["Sudah (HTML)"]
    return [nama, total, c["Sudah (PDF)"], c["Sudah (HTML)"], c["Belum"], sudah / total if total else 0]


def main():
    rows = list(csv.DictReader(open(ROOT / "index.csv", encoding="utf-8")))
    for r in rows:
        r["_status"] = status(r)
    wb, tot = isi(rows)
    out = ROOT / "rekap_download.xlsx"
    wb.save(out)
    print(f"{out}: total {len(rows)}, PDF {tot['Sudah (PDF)']}, HTML {tot['Sudah (HTML)']}, belum {tot['Belum']}")


if __name__ == "__main__":
    main()
